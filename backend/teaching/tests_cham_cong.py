"""BÁO CÁO CHẤM CÔNG theo tháng — V-o, bảng TopHSA dòng 6.3. CHỈ ĐỌC.

Dựng một lớp với buổi đặt ĐÚNG Ở BIÊN: đầu tháng 00:00, đầu tháng sau 00:00, trợ giảng vào
lớp đúng giờ một buổi, rời lớp đúng giờ một buổi khác, điểm danh đúng 24 giờ sau khi hết
buổi (chưa muộn) và 24 giờ + 1 giây (muộn). Tháng 03/2031 — không dữ liệu thật nào nằm đó.

Đi qua VIEW THẬT; CSDL cuộn lại sau mỗi test (`conftest.py`).
"""
import datetime
import io
import uuid

import openpyxl
import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext
from rest_framework.test import APIRequestFactory, force_authenticate

from accounts.models import User
from common.db import q1, x
from common.permissions import ROLE_ACADEMIC, ROLE_ASSISTANT, ROLE_STUDENT, ROLE_TEACHER

f = APIRequestFactory()
pytestmark = pytest.mark.django_db
_luc = datetime.datetime.fromisoformat


def _nguoi(vai, ten):
    r = q1('INSERT INTO users (name, email, password, role, streak) VALUES (%s, %s, %s, %s, 0) RETURNING id',
           (ten, 'cc_%s@example.com' % uuid.uuid4().hex[:10], 'x', vai))
    return User.objects.get(id=r['id'])


def _goi(ai, qs):
    from teaching.cham_cong import ChamCongView
    req = f.get('/x' + qs)
    force_authenticate(req, user=ai)
    return ChamCongView.as_view()(req)


def _buoi(lop, bat_dau, phut, trang_thai, tick_boi=None, tick_luc=None):
    x('INSERT INTO class_sessions (class_id, starts_at, duration_minutes, status, attendance_taken_by, '
      'attendance_taken_at) VALUES (%s, %s, %s, %s, %s, %s)',
      (lop, _luc(bat_dau), phut, trang_thai, tick_boi, _luc(tick_luc) if tick_luc else None))


def _vao(lop, u, tu, den=None):
    x('INSERT INTO class_members (class_id, user_id, joined_at, left_at, leave_reason) VALUES (%s, %s, %s, %s, %s)',
      (lop, u.id, _luc(tu), _luc(den) if den else None, 'dropped' if den else None))


def _lop(gv):
    return q1("INSERT INTO classes (name, teacher_id, status) VALUES (%s, %s, 'active') RETURNING id",
              ('CC lớp %s' % uuid.uuid4().hex[:6], gv.id))['id']


@pytest.fixture
def canh(db):
    d = uuid.uuid4().hex[:6]
    gv = _nguoi(ROLE_TEACHER, 'CC GV %s' % d)
    tg1 = _nguoi(ROLE_ASSISTANT, 'CC TG1 %s' % d)      # ở lớp cả tháng (vào đúng 00:00 ngày 1)
    tg2 = _nguoi(ROLE_ASSISTANT, 'CC TG2 %s' % d)      # vào lớp 15/03
    tg3 = _nguoi(ROLE_ASSISTANT, 'CC TG3 %s' % d)      # rời lớp đúng giờ buổi 05/03
    tg4 = _nguoi(ROLE_ASSISTANT, 'CC TG4 %s' % d)      # không lớp nào
    em = _nguoi(ROLE_STUDENT, 'CC Em %s' % d)
    hv = _nguoi(ROLE_ACADEMIC, 'CC HV %s' % d)
    lop = _lop(gv)
    _vao(lop, tg1, '2031-03-01T00:00')
    _vao(lop, tg2, '2031-03-15T00:00')
    _vao(lop, tg3, '2031-03-01T00:00', '2031-03-05T19:00')
    _vao(lop, em, '2031-02-01T00:00')
    _buoi(lop, '2031-03-01T00:00', 60, 'done')                                             # biên đầu tháng
    _buoi(lop, '2031-03-05T19:00', None, 'done', tg1.id, '2031-03-05T21:00')               # 90' mặc định
    _buoi(lop, '2031-03-12T19:00', 60, 'done', gv.id, '2031-03-13T20:00:01')               # MUỘN (1 giây)
    _buoi(lop, '2031-03-20T19:00', 120, 'planned', gv.id, '2031-03-21T21:00')              # đúng 24h: CHƯA muộn
    _buoi(lop, '2031-03-25T19:00', 60, 'cancelled', gv.id, '2031-03-25T21:00')             # huỷ: không tính
    _buoi(lop, '2031-03-28T19:00', 60, 'planned')                                          # chưa dạy
    _buoi(lop, '2031-04-01T00:00', 60, 'done', gv.id, '2031-04-01T02:00')                  # tháng sau
    return {'gv': gv, 'tg1': tg1, 'tg2': tg2, 'tg3': tg3, 'tg4': tg4, 'em': em, 'hv': hv, 'lop': lop}


def _theo_id(r):
    assert r.status_code == 200, r.data
    return {n['id']: n for n in r.data['nguoi']}


def test_so_buoi_phut_diem_danh_va_muon_tung_nguoi(canh):
    ds = _theo_id(_goi(canh['hv'], '?thang=2031-03'))
    gv, tg1, tg2, tg3, tg4 = (ds[canh[k].id] for k in ('gv', 'tg1', 'tg2', 'tg3', 'tg4'))
    assert (gv['soBuoi'], gv['soPhut'], gv['daDiemDanh'], gv['diemDanhMuon']) == (4, 330, 2, 1), gv
    assert (tg1['soBuoi'], tg1['soPhut'], tg1['daDiemDanh'], tg1['diemDanhMuon']) == (4, 330, 1, 0), tg1
    assert (tg2['soBuoi'], tg2['soPhut']) == (1, 120), tg2
    assert (tg3['soBuoi'], tg3['soPhut']) == (1, 60), tg3
    assert (tg4['soBuoi'], tg4['soPhut'], tg4['daDiemDanh']) == (0, 0, 0), 'trợ giảng không lớp vẫn có dòng'
    assert gv['vai'] == 'Giảng viên' and tg1['vai'] == 'Trợ giảng'
    assert canh['em'].id not in ds and canh['hv'].id not in ds, 'học viên / học vụ không có dòng chấm công'
    r = _goi(canh['hv'], '?thang=2031-03')
    assert r.data['tu'] == '2031-03-01' and r.data['den'] == '2031-03-31' and r.data['lateHours'] == 24
    # Tháng NGẮN: tháng 2 dừng đúng 28/02 — buổi 00:00 ngày 01/03 không lọt sang.
    r = _goi(canh['hv'], '?thang=2031-02')
    assert r.data['den'] == '2031-02-28', r.data['den']
    assert _theo_id(r)[canh['gv'].id]['soBuoi'] == 0


def test_tro_giang_dung_lop_va_o_trong_lop_chi_tinh_mot_lan(canh):
    """Trợ giảng được đặt làm giảng viên chủ lớp VÀ vẫn là thành viên lớp ấy: mỗi buổi
    một lần, không hai."""
    tg = _nguoi(ROLE_ASSISTANT, 'CC TG đứng lớp')
    lop = _lop(tg)
    _vao(lop, tg, '2031-01-01T00:00')
    _buoi(lop, '2031-03-03T08:00', 90, 'done')
    _buoi(lop, '2031-03-10T08:00', 90, 'done')
    n = _theo_id(_goi(canh['hv'], '?thang=2031-03'))[tg.id]
    assert (n['soBuoi'], n['soPhut']) == (2, 180), n


def test_mot_cau_SQL_bat_ke_so_lop_va_so_buoi(canh):
    with CaptureQueriesContext(connection) as it:
        _goi(canh['hv'], '?thang=2031-03')
    gv2 = _nguoi(ROLE_TEACHER, 'CC GV2')
    for _ in range(3):
        lop = _lop(gv2)
        _vao(lop, _nguoi(ROLE_ASSISTANT, 'CC TGx'), '2031-01-01T00:00')
        for ngay in ('2031-03-02', '2031-03-09', '2031-03-16'):
            _buoi(lop, ngay + 'T08:00', 90, 'done')
    with CaptureQueriesContext(connection) as nhieu:
        r = _goi(canh['hv'], '?thang=2031-03')
    assert _theo_id(r)[gv2.id]['soBuoi'] == 9
    assert len(nhieu.captured_queries) == len(it.captured_queries) == 1, \
        [c['sql'][:80] for c in nhieu.captured_queries]


def test_xuat_excel_va_thang_sai(canh):
    r = _goi(canh['hv'], '?thang=2031-03&dinh_dang=xlsx')
    assert r['Content-Type'].startswith('application/vnd.openxmlformats'), r['Content-Type']
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).worksheets[0]
    dau = [c.value for c in ws[1]]
    assert dau[:4] == ['Họ tên', 'Email', 'Vai trò', 'Số buổi đã dạy'], dau
    hang = {row[0].value: [c.value for c in row] for row in ws.iter_rows(min_row=2)}
    gv = hang[canh['gv'].name]
    assert gv[3] == 4 and gv[4] == 330 and gv[5] == 5.5, gv
    for sai in ('2031-13', 'thang-ba', '2031'):
        assert _goi(canh['hv'], '?thang=' + sai).status_code == 400, sai
    assert _goi(canh['hv'], '?thang=2031-03&dinh_dang=pdf').status_code == 400


def test_giang_vien_khong_xem_duoc(canh):
    assert _goi(canh['gv'], '?thang=2031-03').status_code == 403
    assert _goi(canh['tg1'], '?thang=2031-03').status_code == 403
