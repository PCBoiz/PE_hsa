"""HỒ SƠ HỌC VIÊN: Tỉnh/Thành phố, tình trạng học tập, tình trạng học phí — V-m, bảng TopHSA
dòng 3 + 7.

  · Tỉnh/Thành chọn từ 34 đơn vị (`teaching/tinh_thanh.py`); giá trị CŨ gõ tay vẫn hiện
    nguyên văn và vẫn "giữ nguyên" được khi sửa ô khác — chỉ không GHI MỚI được ngoài danh sách.
  · Tình trạng học tập TÍNH từ lượt học (`teaching/tinh_trang.py`) — cùng một biểu thức ở
    hồ sơ, danh sách tài khoản (cột + ô lọc) và tệp xuất.
  · Tình trạng học phí: MỘT ô chọn tay (§63), học vụ / quản trị viên đặt, có nhật ký.

Đi qua VIEW THẬT; CSDL cuộn lại sau mỗi test (`conftest.py`).
"""
import io
import json
import uuid

import openpyxl
import pytest
from rest_framework.test import APIRequestFactory, force_authenticate

from accounts.models import User
from common.db import q1, x
from common.permissions import ROLE_ACADEMIC, ROLE_ADMIN, ROLE_STUDENT, ROLE_TEACHER

f = APIRequestFactory()
pytestmark = pytest.mark.django_db


def _goi(view, method, body=None, ai=None, qs='', **kw):
    req = (getattr(f, method)('/x' + qs, body, format='json') if body is not None
           else getattr(f, method)('/x' + qs))
    force_authenticate(req, user=ai)
    return view.as_view()(req, **kw)


def _nguoi(vai, ten=None):
    r = q1('INSERT INTO users (name, email, password, role, streak) VALUES (%s, %s, %s, %s, 0) RETURNING id',
           (ten or 'TT %s' % vai, 'tt_%s@example.com' % uuid.uuid4().hex[:10], 'x', vai))
    return User.objects.get(id=r['id'])


def _lop(status):
    return q1('INSERT INTO classes (name, status) VALUES (%s, %s) RETURNING id',
              ('TT lớp %s' % uuid.uuid4().hex[:6], status))['id']


def _vao(lop, u, vao='2026-01-01', roi=None, ly_do=None):
    x('INSERT INTO class_members (class_id, user_id, joined_at, left_at, leave_reason) VALUES (%s, %s, %s, %s, %s)',
      (lop, u.id, vao, roi, ly_do))


# ── Học phí: vocab = CHECK ─────────────────────────────────────────────────

def test_hoc_phi_khop_rang_buoc_CSDL():
    from teaching.tinh_trang import MA_HOC_PHI
    r = q1("SELECT pg_get_constraintdef(oid) AS d FROM pg_constraint WHERE conname = 'users_tuition_status_check'")
    assert r, 'chưa có CHECK users_tuition_status_check (§63)'
    import re
    assert set(re.findall(r"'([a-z_]+)'", r['d'])) == set(MA_HOC_PHI), r['d']


def test_hoc_vu_dat_hoc_phi_co_nhat_ky_va_giang_vien_khong_dat_duoc():
    from teaching.ho_so import HoSoHocVienView, MucTieuHocVienView
    hv, em, gv = _nguoi(ROLE_ACADEMIC), _nguoi(ROLE_STUDENT), _nguoi(ROLE_TEACHER)
    r = _goi(HoSoHocVienView, 'patch', {'tuitionStatus': 'sap_het'}, ai=hv, user_id=em.id)
    assert r.status_code == 200, r.data
    assert r.data['profile']['tuitionStatus'] == 'sap_het'
    assert q1('SELECT tuition_status FROM users WHERE id=%s', (em.id,))['tuition_status'] == 'sap_het'
    nk = q1("SELECT detail FROM admin_audit WHERE action='user.profile' AND target_id=%s ORDER BY id DESC LIMIT 1",
            (str(em.id),))
    ct = nk['detail'] if isinstance(nk['detail'], dict) else json.loads(nk['detail'])
    assert ct['cu'] == {'tuition_status': None} and ct['moi'] == {'tuition_status': 'sap_het'}, ct
    r = _goi(HoSoHocVienView, 'patch', {'tuitionStatus': 'no_hoc_phi'}, ai=hv, user_id=em.id)
    assert r.status_code == 400 and 'tuitionStatus' in r.data['errors'], r.data
    assert _goi(HoSoHocVienView, 'patch', {'tuitionStatus': ''}, ai=hv, user_id=em.id).status_code == 200
    assert q1('SELECT tuition_status FROM users WHERE id=%s', (em.id,))['tuition_status'] is None
    # Giảng viên: không vào được hồ sơ đầy đủ; đường mục tiêu của lớp bỏ qua ô học phí.
    assert _goi(HoSoHocVienView, 'patch', {'tuitionStatus': 'het'}, ai=gv, user_id=em.id).status_code == 403
    lop = q1("INSERT INTO classes (name, teacher_id, status) VALUES ('TT gv', %s, 'active') RETURNING id",
             (gv.id,))['id']
    _vao(lop, em)
    _goi(MucTieuHocVienView, 'patch', {'tuitionStatus': 'het', 'studyGoal': 'x'}, ai=gv, class_id=lop, user_id=em.id)
    assert q1('SELECT tuition_status FROM users WHERE id=%s', (em.id,))['tuition_status'] is None


# ── Tỉnh/Thành phố ───────────────────────────────────────────────────────────

def test_tinh_thanh_chon_tu_danh_sach_gia_tri_cu_van_giu():
    from teaching.ho_so import HoSoHocVienView
    hv, em = _nguoi(ROLE_ACADEMIC), _nguoi(ROLE_STUDENT)
    assert _goi(HoSoHocVienView, 'patch', {'region': 'Huế'}, ai=hv, user_id=em.id).status_code == 200
    r = _goi(HoSoHocVienView, 'patch', {'region': 'HN'}, ai=hv, user_id=em.id)
    assert r.status_code == 400 and 'region' in r.data['errors'], r.data
    # Giá trị cũ gõ tay (trước danh sách): đọc ra nguyên văn, gửi lại nguyên văn khi sửa ô khác vẫn được.
    x("UPDATE users SET region = 'Ha Noi (ghi tay)' WHERE id=%s", (em.id,))
    r = _goi(HoSoHocVienView, 'get', ai=hv, user_id=em.id)
    assert r.data['profile']['region'] == 'Ha Noi (ghi tay)'
    r = _goi(HoSoHocVienView, 'patch', {'region': 'Ha Noi (ghi tay)', 'school': 'THPT A'}, ai=hv, user_id=em.id)
    assert r.status_code == 200, r.data
    assert q1('SELECT region FROM users WHERE id=%s', (em.id,))['region'] == 'Ha Noi (ghi tay)'
    assert _goi(HoSoHocVienView, 'patch', {'region': 'Ha Noi 2'}, ai=hv, user_id=em.id).status_code == 400
    assert _goi(HoSoHocVienView, 'patch', {'region': ''}, ai=hv, user_id=em.id).status_code == 200
    assert q1('SELECT region FROM users WHERE id=%s', (em.id,))['region'] is None


# ── Tình trạng học tập (tính) ────────────────────────────────────────────────

def test_tinh_trang_hoc_tinh_tu_luot_hoc_o_ho_so_danh_sach_va_tep_xuat():
    from teaching.admin_users import AdminUsersView
    from teaching.exports import AdminUsersCsvView
    from teaching.ho_so import HoSoHocVienView
    ad = _nguoi(ROLE_ADMIN)
    d = uuid.uuid4().hex[:6]
    em = {k: _nguoi(ROLE_STUDENT, 'TT %s %s' % (k, d)) for k in
          ('dang', 'chua', 'xong', 'xong_lop_ket_thuc', 'nghi', 'xong_roi_nghi', 'chi_lop_huy', 'dang_va_nghi')}
    _vao(_lop('active'), em['dang'])
    _vao(_lop('active'), em['xong'], roi='2026-05-01', ly_do='completed')
    _vao(_lop('finished'), em['xong_lop_ket_thuc'])
    _vao(_lop('active'), em['nghi'], roi='2026-05-01', ly_do='dropped')
    _vao(_lop('finished'), em['xong_roi_nghi'], roi='2026-03-01', ly_do='completed')
    _vao(_lop('active'), em['xong_roi_nghi'], vao='2026-04-01', roi='2026-06-01', ly_do='dropped')
    _vao(_lop('cancelled'), em['chi_lop_huy'])
    _vao(_lop('active'), em['dang_va_nghi'], roi='2026-05-01', ly_do='dropped')
    _vao(_lop('active'), em['dang_va_nghi'], vao='2026-06-01')
    mong = {'dang': 'dang_hoc', 'chua': 'chua_xep_lop', 'xong': 'da_hoc_xong', 'xong_lop_ket_thuc': 'da_hoc_xong',
            'nghi': 'da_nghi', 'xong_roi_nghi': 'da_nghi', 'chi_lop_huy': 'chua_xep_lop', 'dang_va_nghi': 'dang_hoc'}
    # Bảo lưu (lý do rời `reserved`) KHÁC bỏ học — và lượt GẦN NHẤT quyết định.
    em['bao_luu'] = _nguoi(ROLE_STUDENT, 'TT bl %s' % d)
    _vao(_lop('active'), em['bao_luu'], vao='2026-01-01', roi='2026-03-01', ly_do='dropped')
    _vao(_lop('active'), em['bao_luu'], vao='2026-04-01', roi='2026-06-01', ly_do='reserved')
    mong['bao_luu'] = 'bao_luu'
    em['dung'] = _nguoi(ROLE_STUDENT, 'TT dung %s' % d)
    _vao(_lop('paused'), em['dung'])
    _vao(_lop('active'), em['dung'], roi='2026-05-01', ly_do='completed')
    mong['dung'] = 'tam_dung'
    for k, u in em.items():
        r = _goi(HoSoHocVienView, 'get', ai=ad, user_id=u.id)
        assert r.data['profile']['tinhTrangHoc'] == mong[k], (k, r.data['profile']['tinhTrangHoc'])
    gv = _nguoi(ROLE_TEACHER)
    assert _goi(HoSoHocVienView, 'get', ai=ad, user_id=gv.id).data['profile']['tinhTrangHoc'] is None

    ds = _goi(AdminUsersView, 'get', ai=ad, qs='?per_page=100&q=' + d).data
    assert {u['id']: u['tinhTrangHoc'] for u in ds['users']} == {em[k].id: v for k, v in mong.items()}
    assert [o['ma'] for o in ds['tinhTrangHocOptions']][:2] == ['dang_hoc', 'tam_dung']
    loc = _goi(AdminUsersView, 'get', ai=ad, qs='?per_page=100&tinh_trang_hoc=da_nghi&q=' + d).data
    assert {u['id'] for u in loc['users']} == {em['nghi'].id, em['xong_roi_nghi'].id}
    assert _goi(AdminUsersView, 'get', ai=ad, qs='?tinh_trang_hoc=abc&q=' + d).data['total'] == 0

    r = _goi(AdminUsersCsvView, 'get', ai=ad, qs='?dinh_dang=xlsx&q=' + d)
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).worksheets[0]
    dau = [c.value for c in ws[1]]
    assert 'Tình trạng học tập' in dau and 'Học phí' in dau, dau
    hang = {row[0].value: row[dau.index('Tình trạng học tập')].value for row in ws.iter_rows(min_row=2)}
    assert hang[em['xong'].name] == 'Đã học xong' and hang[em['chua'].name] == 'Chưa xếp lớp', hang


def test_loc_hoc_phi_o_danh_sach():
    from teaching.admin_users import AdminUsersView
    hv = _nguoi(ROLE_ACADEMIC)
    d = uuid.uuid4().hex[:6]
    a, b = _nguoi(ROLE_STUDENT, 'TT hp a %s' % d), _nguoi(ROLE_STUDENT, 'TT hp b %s' % d)
    _nguoi(ROLE_TEACHER, 'TT hp gv %s' % d)   # nhân sự: học phí trống nhưng KHÔNG phải "chưa đặt"
    x("UPDATE users SET tuition_status = 'het' WHERE id=%s", (a.id,))
    r = _goi(AdminUsersView, 'get', ai=hv, qs='?hoc_phi=het&q=' + d).data
    assert [u['id'] for u in r['users']] == [a.id] and r['users'][0]['hocPhi'] == 'het'
    assert [o['ma'] for o in r['hocPhiOptions']] == ['da_dong', 'sap_het', 'het', 'bao_luu']
    r = _goi(AdminUsersView, 'get', ai=hv, qs='?hoc_phi=chua_dat&q=' + d).data
    assert [u['id'] for u in r['users']] == [b.id], 'ô "chưa đặt" = học phí trống'
    # Quản trị viên thấy cả nhân sự — "chưa đặt" vẫn chỉ là học viên.
    r = _goi(AdminUsersView, 'get', ai=_nguoi(ROLE_ADMIN), qs='?hoc_phi=chua_dat&q=' + d).data
    assert [u['id'] for u in r['users']] == [b.id], r['users']


# ── Lớp trung tâm + môn đã mở trên hồ sơ (giữ từ b3316cb) ────────────────────

def test_ho_so_si_so_chi_dem_hoc_vien_va_mon_da_mo_tru_khoa_nhap():
    """Sĩ số THẬT của lớp = học viên đang học (trợ giảng trong lớp không tính — cùng luật
    trần lớp gia sư). "Môn đã mở" = đúng thứ cổng mở môn mở cho em: khoá nháp không có."""
    from teaching.ho_so import HoSoHocVienView
    ad = _nguoi(ROLE_ADMIN)
    em = _nguoi(ROLE_STUDENT)
    lop = q1("INSERT INTO classes (name, course_id, status) VALUES (%s, 'hsa_verbal', 'active') RETURNING id",
             ('TT sĩ số %s' % uuid.uuid4().hex[:6],))['id']
    _vao(lop, em)
    _vao(lop, _nguoi('Trợ giảng'))
    nhap = 'tt_nhap_%s' % uuid.uuid4().hex[:6]
    x("INSERT INTO courses (id, title, is_published) VALUES (%s, 'TT khoá nháp', FALSE)", (nhap,))
    lop2 = q1("INSERT INTO classes (name, course_id, status) VALUES (%s, %s, 'active') RETURNING id",
              ('TT nháp %s' % uuid.uuid4().hex[:6], nhap))['id']
    _vao(lop2, em)
    hs = _goi(HoSoHocVienView, 'get', ai=ad, user_id=em.id).data['profile']
    so = {c['id']: c['siSo'] for c in hs['classes']}
    assert so[lop] == 1, 'trợ giảng bị đếm vào sĩ số: %r' % so
    assert {c['id'] for c in hs['enrolledCourses']} == {'hsa_verbal'}, hs['enrolledCourses']
