"""XUẤT EXCEL (.xlsx) + BỘ LỌC trên các bản xuất sẵn có — V-k, bảng TopHSA dòng 6.

Bảng khách đòi "Xuất Excel / CSV" và "bộ lọc thời gian / lớp / môn / khoá". Trước 25/09/2026
chỉ có CSV, không lọc. Luật canh ở đây:

  · MỘT bộ ghi .xlsx (`common/bangtinh.ghi_xlsx`) — ô chữ bắt đầu bằng = + - @ là CHỮ, không
    bao giờ thành công thức (Excel mở ra không chạy gì); số là số (lọc/cộng được), ngày là ngày.
  · Bản xlsx đi qua CÙNG hàng rào với CSV: trợ giảng không nhận cột liên lạc.
  · `tu`/`den` lọc buổi của bảng điểm danh — cột VÀ số tổng cùng khoảng ấy.
  · Danh sách tài khoản lọc thêm theo đợt học, môn, ngày cấp — cùng hàm lọc với màn hình.

Đi qua VIEW THẬT; CSDL cuộn lại sau mỗi test (`conftest.py`).
"""
import datetime
import io
import uuid

import openpyxl
import pytest
from rest_framework.test import APIRequestFactory, force_authenticate

from accounts.models import User
from common.clock import local_now
from common.db import q1, x
from common.permissions import ROLE_ADMIN, ROLE_ASSISTANT, ROLE_STUDENT, ROLE_TEACHER

f = APIRequestFactory()
_luc = datetime.datetime.fromisoformat   # giờ VN naive — quy ước `common/clock.py`
#: Số NGẪU NHIÊN mỗi lượt — `idx_users_phone` là chỉ mục duy nhất, CSDL dev đã có số cố định.
_SDT = '09%08d' % (uuid.uuid4().int % 10**8)
pytestmark = pytest.mark.django_db


def _goi(view, ai, qs='', **kw):
    req = f.get('/x' + qs)
    force_authenticate(req, user=ai)
    return view.as_view()(req, **kw)


def _nguoi(vai, ten=None, phone=None):
    r = q1('INSERT INTO users (name, email, phone, password, role, streak) '
           'VALUES (%s, %s, %s, %s, %s, 0) RETURNING id',
           (ten or 'XE %s' % vai, 'xe_%s@example.com' % uuid.uuid4().hex[:10], phone, 'x', vai))
    return User.objects.get(id=r['id'])


def _bang(resp):
    assert resp['Content-Type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'), resp['Content-Type']
    assert '.xlsx' in resp['Content-Disposition'], resp['Content-Disposition']
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    return wb.worksheets[0]


@pytest.fixture
def lop(db):
    gv = _nguoi(ROLE_TEACHER)
    em = _nguoi(ROLE_STUDENT, ten='=HYPERLINK("http://x","bam") %s' % uuid.uuid4().hex[:4], phone=_SDT)
    c = q1("INSERT INTO classes (name, course_id, teacher_id, status) "
           "VALUES (%s, 'hsa_quantitative', %s, 'active') RETURNING id",
           ('XE lớp %s' % uuid.uuid4().hex[:6], gv.id))['id']
    x('INSERT INTO class_members (class_id, user_id, joined_at) VALUES (%s, %s, %s)',
      (c, em.id, _luc('2026-08-01T00:00')))
    # Hai buổi đã dạy ở HAI tháng khác nhau: 10/08 em có mặt, 10/09 em vắng.
    ids = []
    for ngay, tt in ((_luc('2026-08-10T19:00'), 'present'),
                     (_luc('2026-09-10T19:00'), 'absent')):
        s = q1("INSERT INTO class_sessions (class_id, starts_at, duration_minutes, status) "
               "VALUES (%s, %s, 90, 'done') RETURNING id", (c, ngay))['id']
        x('INSERT INTO attendance (session_id, user_id, status) VALUES (%s, %s, %s)', (s, em.id, tt))
        ids.append(s)
    return {'id': c, 'gv': gv, 'em': em, 'buoi': ids}


# ── Bộ ghi .xlsx dùng chung ──────────────────────────────────────────────────

def test_ghi_xlsx_o_cong_thuc_la_chu_so_la_so_ngay_la_ngay():
    from common.bangtinh import Trang, ghi_xlsx
    du_lieu = ghi_xlsx([Trang('Thử', [
        ['Chữ', 'Số', 'Ngày', 'Rác'],
        ['=1+1', 42, _luc('2026-09-25T07:30'), 'a\x00b\x07c'],
        ['@SUM(A1)', 3.5, datetime.date(2026, 9, 1), '-2+3'],
    ])])
    wb = openpyxl.load_workbook(io.BytesIO(du_lieu))    # data_only=False: thấy công thức nếu có
    ws = wb['Thử']
    for o in (ws['A2'], ws['A3'], ws['D3']):
        assert o.data_type == 's', (o.coordinate, o.data_type, o.value)
        assert o.quotePrefix, 'ô chữ trông như công thức phải mang dấu nháy ẩn (%s)' % o.coordinate
    assert ws['A2'].value == '=1+1'
    assert ws['B2'].value == 42 and ws['B3'].value == 3.5
    assert isinstance(ws['C2'].value, datetime.datetime)
    assert ws['D2'].value == 'abc', 'ký tự điều khiển phải bị bỏ (openpyxl ném lỗi → 500)'
    assert ws['A1'].font.b, 'dòng tiêu đề in đậm'


# ── Bản xlsx của các tệp lớp ─────────────────────────────────────────────────

def test_tro_giang_tai_xlsx_KHONG_co_cot_lien_lac(lop):
    from teaching.exports import ClassAttendanceCsvView, ClassProgressCsvView
    tg = _nguoi(ROLE_ASSISTANT)
    x('INSERT INTO class_members (class_id, user_id, joined_at) VALUES (%s, %s, now())', (lop['id'], tg.id))
    for view in (ClassProgressCsvView, ClassAttendanceCsvView):
        ws = _bang(_goi(view, tg, '?dinh_dang=xlsx', class_id=lop['id']))
        dau = [c.value for c in ws[1]]
        assert 'Email' not in dau and 'Số điện thoại' not in dau, (view.__name__, dau)
        than = ' '.join(str(c.value) for hang in ws.iter_rows() for c in hang if c.value is not None)
        assert lop['em'].email not in than and _SDT not in than, view.__name__


def test_giang_vien_tai_xlsx_tien_do_du_cot_so_dien_thoai_giu_so_0(lop):
    from teaching.exports import ClassProgressCsvView
    ws = _bang(_goi(ClassProgressCsvView, lop['gv'], '?dinh_dang=xlsx', class_id=lop['id']))
    dau = [c.value for c in ws[1]]
    assert 'Email' in dau and 'Số điện thoại' in dau
    dong = [c for c in ws[2]]
    o_sdt = dong[dau.index('Số điện thoại')]
    assert o_sdt.value == _SDT and o_sdt.data_type == 's', (o_sdt.value, o_sdt.data_type)
    o_ten = dong[dau.index('Họ tên')]
    assert o_ten.data_type == 's' and o_ten.value.startswith('=HYPERLINK'), (o_ten.data_type, o_ten.value)
    # CSV vẫn là mặc định — không gửi `dinh_dang` thì ra đúng tệp cũ.
    r = _goi(ClassProgressCsvView, lop['gv'], class_id=lop['id'])
    assert r['Content-Type'].startswith('text/csv')
    assert _goi(ClassProgressCsvView, lop['gv'], '?dinh_dang=pdf', class_id=lop['id']).status_code == 400


def test_diem_danh_loc_tu_den_cot_va_so_tong_cung_khoang(lop):
    from teaching.exports import ClassAttendanceCsvView
    ws = _bang(_goi(ClassAttendanceCsvView, lop['gv'], '?dinh_dang=xlsx&tu=2026-09-01&den=2026-09-30',
                    class_id=lop['id']))
    dau = [c.value for c in ws[1]]
    buoi = [h for h in dau if isinstance(h, str) and h[:2].isdigit()]
    assert buoi == ['10/09 19:00'], dau
    # `den` tính CẢ ngày: buổi 19:00 ngày 10/09 nằm trong khoảng 10/09 – 10/09.
    ws1 = _bang(_goi(ClassAttendanceCsvView, lop['gv'], '?dinh_dang=xlsx&tu=2026-09-10&den=2026-09-10',
                     class_id=lop['id']))
    assert [h for h in (c.value for c in ws1[1]) if isinstance(h, str) and h[:2].isdigit()] == ['10/09 19:00']
    hang = {dau[i]: c.value for i, c in enumerate(ws[2])}
    assert hang['Có mặt'] == 0 and hang['Vắng'] == 1, hang
    # Không lọc: đủ hai buổi, tổng đủ hai buổi.
    ws = _bang(_goi(ClassAttendanceCsvView, lop['gv'], '?dinh_dang=xlsx', class_id=lop['id']))
    dau = [c.value for c in ws[1]]
    hang = {dau[i]: c.value for i, c in enumerate(ws[2])}
    assert hang['Có mặt'] == 1 and hang['Vắng'] == 1, hang
    r = _goi(ClassAttendanceCsvView, lop['gv'], '?tu=hom-qua', class_id=lop['id'])
    assert r.status_code == 400
    r = _goi(ClassAttendanceCsvView, lop['gv'], '?tu=2026-09-30&den=2026-09-01', class_id=lop['id'])
    assert r.status_code == 400


# ── Danh sách tài khoản: lọc theo đợt, môn, ngày cấp ───────────────────────

def test_tai_khoan_loc_theo_dot_mon_ngay_cap(lop):
    from teaching.exports import AdminUsersCsvView
    ad = _nguoi(ROLE_ADMIN)
    dot = q1("INSERT INTO terms (name, starts_on, ends_on, status) "
             "VALUES (%s, '2026-08-01', '2026-12-31', 'active') RETURNING id",
             ('XE đợt %s' % uuid.uuid4().hex[:6],))['id']
    x('UPDATE classes SET term_id = %s WHERE id = %s', (dot, lop['id']))
    duoi = uuid.uuid4().hex[:6]
    # `khac`: lớp môn Định tính, KHÔNG thuộc đợt nào. `ca_ba`: lớp để trống môn (học cả ba
    # môn HSA), thuộc đợt. Tài khoản `khac` cấp ngày 05/01/2025.
    khac = _nguoi(ROLE_STUDENT, ten='XE khác %s' % duoi)
    ca_ba = _nguoi(ROLE_STUDENT, ten='XE ca ba %s' % duoi)
    for em, mon, d in ((khac, 'hsa_verbal', None), (ca_ba, None, dot)):
        c = q1("INSERT INTO classes (name, course_id, term_id, status) VALUES (%s, %s, %s, 'active') "
               'RETURNING id', ('XE lớp %s' % uuid.uuid4().hex[:6], mon, d))['id']
        x('INSERT INTO class_members (class_id, user_id, joined_at) VALUES (%s, %s, now())', (c, em.id))
    x("UPDATE users SET created_at = %s WHERE id = %s", (_luc('2025-01-05T00:00'), khac.id))
    # Mã cấp tài khoản ghi `created_at = local_now()` (giờ VN); `DEFAULT now()` của CSDL là
    # giờ UTC — lệch ngày lúc 0h–7h sáng. Đặt như mã thật làm.
    x('UPDATE users SET created_at = %s WHERE id = ANY(%s)', (local_now(), [lop['em'].id, ca_ba.id]))

    def ten(qs):
        ws = _bang(_goi(AdminUsersCsvView, ad, '?dinh_dang=xlsx&q=' + duoi + '&' + qs))
        return {r[0].value for r in ws.iter_rows(min_row=2)} - {None}

    # `q=<đuôi>` khoanh về đúng các em của phép kiểm này (tên mang đuôi ngẫu nhiên).
    x("UPDATE users SET name = name || ' ' || %s WHERE id = %s", (duoi, lop['em'].id))
    em = q1('SELECT name FROM users WHERE id = %s', (lop['em'].id,))['name']
    assert ten('term_id=%d' % dot) == {em, ca_ba.name}
    assert ten('course_id=hsa_quantitative') == {em, ca_ba.name}
    assert ten('course_id=hsa_verbal') == {khac.name, ca_ba.name}
    assert ten('course_id=hsa_verbal&term_id=%d' % dot) == {ca_ba.name}
    assert ten('term_id=abc') == set()
    # Ngày cấp tài khoản: `den` tính CẢ ngày ấy.
    hom_nay = local_now().date().isoformat()
    assert ten('tu=%s&den=%s' % (hom_nay, hom_nay)) == {em, ca_ba.name}
    assert ten('tu=2025-01-05&den=2025-01-05') == {khac.name}
    assert _goi(AdminUsersCsvView, ad, '?tu=05/01/2025').status_code == 400
