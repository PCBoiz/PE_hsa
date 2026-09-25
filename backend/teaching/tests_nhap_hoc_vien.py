"""NHẬP HỌC VIÊN VÀO LỚP TỪ TỆP MẪU — V-j, bảng TopHSA dòng 4 ("import DS theo biểu mẫu").

Học vụ tải tệp mẫu .xlsx, điền, tải lên: XEM TRƯỚC báo từng dòng (tạo tài khoản mới / thêm
tài khoản có sẵn vào lớp / đã trong lớp / lỗi kèm lý do), rồi mới nhập thật.

Luật canh ở đây:
  · khớp tài khoản có sẵn theo mã học viên / email / số điện thoại — ba cách trỏ hai người
    khác nhau thì BÁO, không đoán;
  · tài khoản MỚI đi qua `admin_users.cap_tai_khoan` — CÙNG một đường với ô dán ở trang
    Tài khoản (mã HSA, mật khẩu tạm, nhật ký), không có bản thứ hai;
  · tài khoản CÓ SẴN vào lớp qua `AdminClassMembersView._ghi_thanh_vien` (nhật ký, trần gia
    sư, quên đệm quyền);
  · trần 50 dòng một tệp; lớp gia sư tối đa 3 em — từ chối TRƯỚC khi tạo gì;
  · ô công thức: tệp .xlsx có công thức chưa tính → cả tệp bị từ chối; họ tên mở đầu bằng
    = + - @ (dán từ CSV) → dòng ấy lỗi, không vào CSDL.

Đi qua VIEW THẬT; CSDL cuộn lại sau mỗi test (`conftest.py`).
"""
import io
import json
import uuid

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIRequestFactory, force_authenticate

from accounts.models import User
from common.db import q, q1, x
from common.permissions import ROLE_ACADEMIC, ROLE_STUDENT, ROLE_TEACHER

f = APIRequestFactory()
pytestmark = pytest.mark.django_db
TIEU_DE = ['Họ và tên', 'Email', 'Số điện thoại', 'Mã học viên']


def _sdt():
    return '09%08d' % (uuid.uuid4().int % 10**8)


def _nguoi(vai, **kw):
    r = q1('INSERT INTO users (name, email, phone, password, role, streak) '
           'VALUES (%s, %s, %s, %s, %s, 0) RETURNING id',
           (kw.get('ten', 'NH %s' % vai), kw.get('email') or 'nh_%s@example.com' % uuid.uuid4().hex[:10],
            kw.get('phone'), 'x', vai))
    return User.objects.get(id=r['id'])


def _xlsx(dong, cong_thuc=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(TIEU_DE)
    for d in dong:
        ws.append(d)
    # openpyxl coi MỌI chuỗi mở đầu bằng "=" là công thức. Ô người dùng gõ chữ "=…" trong Excel
    # (có dấu nháy) lưu thành CHỮ — dựng đúng như thế; ô công thức thật đi qua `cong_thuc`.
    for hang in ws.iter_rows(min_row=2):
        for o in hang:
            if isinstance(o.value, str) and o.value.startswith('='):
                o.data_type = 's'
    if cong_thuc:
        ws[cong_thuc[0]] = cong_thuc[1]       # openpyxl lưu CÔNG THỨC, không lưu kết quả
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _nhap(canh, noi_dung, ten='ds.xlsx', dry=True, ai=None, lop=None, **them):
    from teaching.nhap_hoc_vien import NhapHocVienView
    than = {'tep': SimpleUploadedFile(ten, noi_dung), 'dry_run': '1' if dry else '0', **them}
    req = f.post('/x', than, format='multipart')
    force_authenticate(req, user=ai or canh['hocvu'])
    return NhapHocVienView.as_view()(req, class_id=lop or canh['lop'])


@pytest.fixture
def canh(db):
    from teaching.ho_so import cap_ma_hoc_vien
    duoi = uuid.uuid4().hex[:6]
    hocvu = _nguoi(ROLE_ACADEMIC)
    cu = _nguoi(ROLE_STUDENT, ten='NH Em Cũ %s' % duoi, phone=_sdt())
    ma = cap_ma_hoc_vien(cu.id)
    gv = _nguoi(ROLE_TEACHER, ten='NH GV %s' % duoi)
    lop = q1("INSERT INTO classes (name, course_id, status) VALUES (%s, 'hsa_verbal', 'active') RETURNING id",
             ('NH lớp %s' % duoi,))['id']
    return {'hocvu': hocvu, 'cu': cu, 'ma': ma, 'gv': gv, 'lop': lop, 'duoi': duoi}


def _thanh_vien(lop):
    return {r['user_id'] for r in q('SELECT user_id FROM class_members WHERE class_id=%s AND left_at IS NULL',
                                     (lop,))}


# ── Tệp mẫu ──────────────────────────────────────────────────────────────────

def test_tep_mau_doc_lai_duoc_bang_chinh_bo_doc(canh):
    from common.bangtinh import doc, thanh_ban_ghi
    from teaching.nhap_hoc_vien import COT_BAT_BUOC, TEN_KHAC, TepMauNhapHocVienView
    req = f.get('/x')
    force_authenticate(req, user=canh['hocvu'])
    r = TepMauNhapHocVienView.as_view()(req, class_id=canh['lop'])
    assert r.status_code == 200
    assert r['Content-Type'].startswith('application/vnd.openxmlformats')
    ban = thanh_ban_ghi(doc('mau.xlsx', r.content), COT_BAT_BUOC, TEN_KHAC)
    assert ban and all(b['họ và tên'] for _, b in ban), ban


# ── Xem trước: báo từng dòng, không ghi gì ─────────────────────────────────

def test_xem_truoc_bao_tung_dong_va_khong_ghi_gi(canh):
    d = canh['duoi']
    moi = 'nh_moi_%s@example.com' % d
    tep = _xlsx([
        ['Nguyễn Mới %s' % d, moi, '', ''],                                  # 2: tạo mới
        ['(tên trên phiếu khác)', '', '', canh['ma']],                        # 3: có sẵn theo mã
        ['NH Em Cũ', '', canh['cu'].phone, ''],                               # 4: có sẵn theo SĐT (trùng 3)
        ['Trùng Email %s' % d, moi, '', ''],                                  # 5: trùng dòng 2
        ['GV nhầm', canh['gv'].email, '', ''],                                # 6: tài khoản nhân sự
        ['Sai Số %s' % d, '', '0123', ''],                                    # 7: số sai
        ['=HYPERLINK("http://x","bam")', 'nh_ct_%s@example.com' % d, '', ''],  # 8: họ tên như công thức
        ['Thiếu Liên Hệ %s' % d, '', '', ''],                                 # 9: không email/SĐT/mã
        ['Mã Lạ', '', '', 'HSA-99999'],                                       # 10: mã không có
        ['Hai Người', canh['gv'].email, canh['cu'].phone, ''],                # 11: hai định danh, hai người
    ])
    truoc = q1('SELECT count(*) AS n FROM users')['n']
    r = _nhap(canh, tep)
    assert r.status_code == 200, r.data
    kq = {row['line']: row for row in r.data['rows']}
    assert kq[2]['status'] == 'tao_moi', kq[2]
    assert kq[3]['status'] == 'them_vao_lop' and kq[3]['userId'] == canh['cu'].id, kq[3]
    assert kq[4]['status'] == 'loi' and 'dòng 3' in kq[4]['reason'], kq[4]
    assert kq[5]['status'] == 'loi' and 'dòng 2' in kq[5]['reason'], kq[5]
    assert kq[6]['status'] == 'loi' and 'nhân sự' in kq[6]['reason'], kq[6]
    assert kq[7]['status'] == 'loi', kq[7]
    assert kq[8]['status'] == 'loi' and 'công thức' in kq[8]['reason'], kq[8]
    assert kq[9]['status'] == 'loi', kq[9]
    assert kq[10]['status'] == 'loi' and 'HSA-99999' in kq[10]['reason'], kq[10]
    assert kq[11]['status'] == 'loi' and 'hai tài khoản' in kq[11]['reason'], kq[11]
    assert r.data['dem'] == {'tao_moi': 1, 'them_vao_lop': 1, 'da_trong_lop': 0, 'loi': 8}, r.data['dem']
    assert q1('SELECT count(*) AS n FROM users')['n'] == truoc, 'xem trước KHÔNG được tạo tài khoản'
    assert _thanh_vien(canh['lop']) == set(), 'xem trước KHÔNG được xếp lớp'


# ── Nhập thật ────────────────────────────────────────────────────────────────

def test_nhap_that_tao_moi_them_co_san_va_ghi_nhat_ky(canh):
    from courses.truy_cap import quen_truy_cap, quyen_khoa
    d = canh['duoi']
    moi = 'nh_that_%s@example.com' % d
    x('INSERT INTO class_members (class_id, user_id, joined_at) VALUES (%s, %s, now())',
      (canh['lop'], _nguoi(ROLE_STUDENT).id))                       # một em có sẵn trong lớp
    quen_truy_cap(canh['cu'].id)
    assert 'hsa_verbal' not in quyen_khoa(canh['cu'])               # đệm NÓNG, chưa mở môn
    tep = _xlsx([['Trần Mới %s' % d, moi, _sdt(), ''], ['', '', '', canh['ma']]])
    r = _nhap(canh, tep, dry=False, joined_at='2026-09-01')
    assert r.status_code == 201, r.data
    kq = {row['line']: row for row in r.data['rows']}
    em_moi = q1('SELECT id, student_code, must_change_password FROM users WHERE email=%s', (moi,))
    assert em_moi and em_moi['student_code'].startswith('HSA-') and em_moi['must_change_password'], em_moi
    assert kq[2]['status'] == 'tao_moi' and kq[2]['tempPassword'] and kq[2]['studentCode'] == em_moi['student_code']
    assert kq[3]['status'] == 'them_vao_lop'
    assert {em_moi['id'], canh['cu'].id} <= _thanh_vien(canh['lop'])
    vao = q1('SELECT joined_at FROM class_members WHERE class_id=%s AND user_id=%s AND left_at IS NULL',
             (canh['lop'], em_moi['id']))['joined_at']
    assert vao.date().isoformat() == '2026-09-01', vao
    assert quyen_khoa(canh['cu']).get('hsa_verbal') == 'hoc', 'em có sẵn vào lớp mà đệm quyền chưa quên'
    # Nhật ký: tài khoản mới (user.create), em có sẵn (class.member.add), và MỘT dòng cho lượt nhập.
    nk = q("SELECT action, target_type, target_id, detail FROM admin_audit WHERE actor_id=%s ORDER BY id",
           (canh['hocvu'].id,))
    hanh = [n['action'] for n in nk]
    assert hanh.count('user.create') == 1 and hanh.count('class.member.add') == 1, hanh
    nhap = [n for n in nk if n['action'] == 'class.member.import']
    assert len(nhap) == 1 and nhap[0]['target_type'] == 'class' and nhap[0]['target_id'] == str(canh['lop'])
    ct = nhap[0]['detail'] if isinstance(nhap[0]['detail'], dict) else json.loads(nhap[0]['detail'])
    assert ct['tep'] == 'ds.xlsx' and ct['moi'] == [em_moi['id']] and ct['coSan'] == [canh['cu'].id], ct
    # Nhập lại cùng tệp: không nhân đôi — cả hai đã trong lớp / đã có tài khoản.
    r2 = _nhap(canh, tep, dry=True)
    kq2 = {row['line']: row['status'] for row in r2.data['rows']}
    assert kq2 == {2: 'da_trong_lop', 3: 'da_trong_lop'}, kq2


def test_tran_50_dong_mot_tep(canh):
    d = canh['duoi']
    tep = _xlsx([['Em %d %s' % (i, d), 'nh_50_%s_%d@example.com' % (d, i), '', ''] for i in range(51)])
    r = _nhap(canh, tep)
    assert r.status_code == 400 and '50' in r.data['error'], r.data


def test_lop_gia_su_du_cho_bi_chan_truoc_khi_tao(canh):
    d = canh['duoi']
    x("UPDATE classes SET class_type='gia_su' WHERE id=%s", (canh['lop'],))
    for _ in range(2):
        x('INSERT INTO class_members (class_id, user_id, joined_at) VALUES (%s, %s, now())',
          (canh['lop'], _nguoi(ROLE_STUDENT).id))
    tep = _xlsx([['Gia Su A %s' % d, 'nh_gs_a_%s@example.com' % d, '', ''], ['', '', '', canh['ma']]])
    xem = _nhap(canh, tep)
    assert xem.status_code == 200 and any('gia sư' in w for w in xem.data['warnings']), xem.data
    r = _nhap(canh, tep, dry=False)
    assert r.status_code == 400 and 'gia sư' in r.data['error'], r.data
    assert not q1('SELECT 1 AS c FROM users WHERE email=%s', ('nh_gs_a_%s@example.com' % d,))
    assert canh['cu'].id not in _thanh_vien(canh['lop'])


def test_cong_thuc_trong_tep_bi_vo_hieu(canh):
    d = canh['duoi']
    # .xlsx có ô CÔNG THỨC chưa được Excel tính: cả tệp bị từ chối, nói rõ ô nào.
    tep = _xlsx([['Tên %s' % d, 'nh_f_%s@example.com' % d, '', '']],
                cong_thuc=('A2', '=HYPERLINK("http://x","bam")'))
    r = _nhap(canh, tep)
    assert r.status_code == 400 and 'A2' in r.data['error'] and 'CÔNG THỨC' in r.data['error'], r.data
    # .csv: ô "=..." là chữ trần — dòng ấy lỗi, nhập thật cũng không tạo tài khoản nào.
    csv = ('Họ và tên,Email,Số điện thoại,Mã học viên\r\n'
           '=cmd|\' /C calc\'!A0,nh_csv_%s@example.com,,\r\n' % d).encode('utf-8')
    r = _nhap(canh, csv, ten='ds.csv', dry=False)
    assert r.status_code == 201, r.data
    assert r.data['rows'][0]['status'] == 'loi' and r.data['dem']['tao_moi'] == 0, r.data
    assert not q1('SELECT 1 AS c FROM users WHERE email=%s', ('nh_csv_%s@example.com' % d,))


def test_giang_vien_khong_nhap_duoc(canh):
    r = _nhap(canh, _xlsx([['A', 'a@example.com', '', '']]), ai=canh['gv'])
    assert r.status_code == 403
